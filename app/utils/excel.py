"""Excel 导出辅助（openpyxl）。"""
from io import BytesIO
from urllib.parse import quote

from flask import send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def build_workbook(title, headers, rows, widths=None):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31] or "Sheet1"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="409EFF")
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for r, row in enumerate(rows, start=2):
        for c, value in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=value)

    for idx, header in enumerate(headers, start=1):
        width = (widths or {}).get(header, max(12, len(str(header)) * 2 + 4))
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = "A2"
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def excel_response(buffer, filename):
    """返回带 UTF-8 文件名的 xlsx 附件（与旧系统 content-disposition 兼容）。"""
    if not filename.lower().endswith(".xlsx"):
        filename += ".xlsx"
    response = send_file(
        buffer,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        as_attachment=True,
        download_name=filename,
    )
    response.headers["Content-Disposition"] = (
        "attachment; filename=\"%s\"; filename*=utf-8''%s"
        % (filename.encode("ascii", "ignore").decode() or "export.xlsx", quote(filename))
    )
    return response